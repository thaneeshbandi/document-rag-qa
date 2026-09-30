"""
Document loaders.

Each loader reads one file type and returns plain text.
load_document() picks the right loader based on the file extension and
returns every document in the same simple format:

    {"text": "...", "source": "filename.pdf"}
"""

import csv
import json
import os

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = [".pdf", ".docx", ".csv", ".xlsx", ".txt", ".json"]


def load_pdf(file_path):
    reader = PdfReader(file_path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def load_docx(file_path):
    document = Document(file_path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    return "\n".join(paragraphs)


def rows_to_text(header, rows):
    # Turn each table row into a line like "name: Alice, city: Paris"
    # so the column names stay next to their values.
    lines = []
    for row in rows:
        pairs = [f"{column}: {value}" for column, value in zip(header, row) if value not in (None, "")]
        if pairs:
            lines.append(", ".join(pairs))
    return "\n".join(lines)


def load_csv(file_path):
    with open(file_path, newline="", encoding="utf-8", errors="ignore") as file:
        rows = list(csv.reader(file))
    if not rows:
        return ""
    return rows_to_text(rows[0], rows[1:])


def load_xlsx(file_path):
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet_texts = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if rows:
            sheet_texts.append(f"Sheet: {sheet.title}\n" + rows_to_text(rows[0], rows[1:]))
    workbook.close()
    return "\n\n".join(sheet_texts)


def load_txt(file_path):
    with open(file_path, encoding="utf-8", errors="ignore") as file:
        return file.read()


def json_to_lines(value, prefix=""):
    # Turn nested JSON into readable lines like "faq > answer: Support is available..."
    # (plain sentences embed much better than raw JSON with brackets and quotes).
    lines = []
    if isinstance(value, dict):
        for key, item in value.items():
            lines += json_to_lines(item, f"{prefix}{key} > ")
    elif isinstance(value, list):
        for item in value:
            lines += json_to_lines(item, prefix)
    elif prefix:
        lines.append(f"{prefix[:-3]}: {value}")  # [:-3] removes the trailing " > "
    else:
        lines.append(str(value))
    return lines


def load_json(file_path):
    with open(file_path, encoding="utf-8") as file:
        data = json.load(file)
    return "\n".join(json_to_lines(data))


def is_supported_file(filename):
    extension = os.path.splitext(filename)[1].lower()
    return extension in SUPPORTED_EXTENSIONS


def load_document(file_path):
    """Read any supported file and return {"text": ..., "source": ...}."""
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        text = load_pdf(file_path)
    elif extension == ".docx":
        text = load_docx(file_path)
    elif extension == ".csv":
        text = load_csv(file_path)
    elif extension == ".xlsx":
        text = load_xlsx(file_path)
    elif extension == ".txt":
        text = load_txt(file_path)
    elif extension == ".json":
        text = load_json(file_path)
    else:
        raise ValueError(f"Unsupported file type: {extension}")

    return {"text": text, "source": os.path.basename(file_path)}
